import datetime as dt
import io
from decimal import Decimal

from django.db import transaction
from django.db.models import Count
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import TIPOS_CARGA, Bloque, Ejercicio, EjercicioHecho, PesoEjercicio, RegistroPeso, Rutina, Sesion

VENTANA_MEDIA = 4
LIMITE_DIAS = 366  # la fecha de una sesión puede estar a ±1 año de hoy


def fecha_valida(texto, hoy=None):
    """Fecha ISO enviada por el usuario; si es inválida o absurda, usa hoy."""
    hoy = hoy or timezone.localdate()
    try:
        fecha = dt.date.fromisoformat(texto or '')
    except ValueError:
        return hoy
    return fecha if abs((fecha - hoy).days) <= LIMITE_DIAS else hoy


def rango_valido(desde_txt, hasta_txt, hoy=None):
    """(desde, hasta) ISO enviados por el usuario; inválidos -> semana actual. Si vienen
    invertidos se ordenan."""
    hoy = hoy or timezone.localdate()
    try:
        desde, hasta = dt.date.fromisoformat(desde_txt or ''), dt.date.fromisoformat(hasta_txt or '')
    except ValueError:
        desde = lunes_de(hoy)
        return desde, desde + dt.timedelta(days=6)
    return (desde, hasta) if desde <= hasta else (hasta, desde)


def rutina_activa(usuario):
    return Rutina.objects.filter(usuario=usuario, activa=True).first()


def lunes_de(fecha):
    return fecha - dt.timedelta(days=fecha.weekday())


def dia_de_hoy(rutina, hoy):
    return rutina.dias.filter(dia_semana=hoy.weekday()).first()


def proximo_dia(rutina, hoy):
    """Próximo día de la rutina con día de semana fijo, contando desde mañana."""
    dias = {d.dia_semana: d for d in rutina.dias.filter(dia_semana__isnull=False)}
    for adelante in range(1, 8):
        fecha = hoy + dt.timedelta(days=adelante)
        if fecha.weekday() in dias:
            return dias[fecha.weekday()], fecha
    return None, None


def total_ejercicios(dia):
    return Ejercicio.objects.filter(bloque__dia=dia).count()


def hechos_de(usuario, dia, fecha):
    """IDs de ejercicios marcados para ese día en esa fecha."""
    return set(
        EjercicioHecho.objects.filter(
            sesion__usuario=usuario, sesion__dia=dia, sesion__fecha=fecha
        ).values_list('ejercicio_id', flat=True)
    )


@transaction.atomic
def alternar_ejercicio(usuario, ejercicio, fecha):
    """Marca o desmarca un ejercicio; devuelve (hecho, hechos, total, completada)."""
    dia = ejercicio.bloque.dia
    sesion, _ = Sesion.objects.get_or_create(usuario=usuario, dia=dia, fecha=fecha)
    hecho, creado = EjercicioHecho.objects.get_or_create(sesion=sesion, ejercicio=ejercicio)
    if not creado:
        hecho.delete()
    total = total_ejercicios(dia)
    hechos = sesion.hechos.count()
    completada = total > 0 and hechos >= total
    if sesion.completada != completada:
        sesion.completada = completada
        sesion.save(update_fields=['completada'])
    return creado, hechos, total, completada


def pesos_de(usuario, dia, fecha):
    """{ejercicio_id: PesoEjercicio} registrados para ese día en esa fecha."""
    return {
        p.ejercicio_id: p
        for p in PesoEjercicio.objects.filter(
            sesion__usuario=usuario, sesion__dia=dia, sesion__fecha=fecha
        )
    }


def ultimos_pesos(usuario, dia, fecha):
    """{ejercicio_id: PesoEjercicio} más reciente anterior a `fecha` (como referencia)."""
    ultimos = {}
    registros = PesoEjercicio.objects.filter(
        sesion__usuario=usuario, ejercicio__bloque__dia=dia, sesion__fecha__lt=fecha
    ).order_by('-sesion__fecha')
    for p in registros:
        ultimos.setdefault(p.ejercicio_id, p)
    return ultimos


@transaction.atomic
def guardar_pesos(usuario, dia, fecha, cargas):
    """Guarda {ejercicio_id: Carga | None}; None borra el registro. Devuelve la sesión o None."""
    sesion = Sesion.objects.filter(usuario=usuario, dia=dia, fecha=fecha).first()
    if sesion is None:
        if all(c is None for c in cargas.values()):
            return None  # nada que guardar: no se crea una sesión vacía
        sesion = Sesion.objects.create(usuario=usuario, dia=dia, fecha=fecha)
    for ejercicio_id, carga in cargas.items():
        if carga is None:
            PesoEjercicio.objects.filter(sesion=sesion, ejercicio_id=ejercicio_id).delete()
        else:
            PesoEjercicio.objects.update_or_create(
                sesion=sesion, ejercicio_id=ejercicio_id,
                defaults={
                    'modo': carga.modo, 'kg': carga.kg, 'kg_barra': carga.barra,
                    'kg_discos': carga.discos, 'kg_mancuerna': carga.mancuerna,
                },
            )
    return sesion


def resumen_semana(usuario, rutina, hoy):
    """Días de la rutina completados esta semana vs. días de la rutina."""
    lunes = lunes_de(hoy)
    domingo = lunes + dt.timedelta(days=6)
    hechas = Sesion.objects.filter(
        usuario=usuario, dia__rutina=rutina, fecha__range=(lunes, domingo), completada=True
    ).values('dia').distinct().count()
    return hechas, rutina.dias.count()


def _resumen_sesion(sesion, ejercicios, hechos, pesos):
    """Totales de una sesión: ejercicios hechos y carga total (suma de los pesos anotados)."""
    carga = sum((pesos[(sesion.pk, e.pk)].kg for e in ejercicios if (sesion.pk, e.pk) in pesos), Decimal(0))
    return {
        'fecha': sesion.fecha,
        'completada': sesion.completada,
        'hechos': sum((sesion.pk, e.pk) in hechos for e in ejercicios),
        'total': len(ejercicios),
        'carga': carga,
        'con_peso': sum((sesion.pk, e.pk) in pesos for e in ejercicios),
    }


def cierre_semanal(usuario, rutina, hoy):
    """Comparación día por día: semana anterior (cierre) vs. semana en curso.

    Devuelve {'dias': [...], 'anterior': (lunes, domingo), 'actual': (lunes, domingo),
    'carga_anterior': Decimal, 'carga_actual': Decimal}. Cada día trae su resumen de ambas
    semanas y, por ejercicio con peso, el registro de cada una (para ver total o detalle).
    """
    lunes = lunes_de(hoy)
    lunes_ant = lunes - dt.timedelta(days=7)
    rango_ant = (lunes_ant, lunes - dt.timedelta(days=1))
    rango_act = (lunes, lunes + dt.timedelta(days=6))

    dias = list(rutina.dias.order_by('orden'))
    ejercicios = {d.pk: [] for d in dias}
    for e in Ejercicio.objects.filter(bloque__dia__in=dias).select_related('bloque').order_by('bloque__orden', 'orden'):
        ejercicios[e.bloque.dia_id].append(e)

    sesiones = list(
        Sesion.objects.filter(usuario=usuario, dia__in=dias, fecha__range=(lunes_ant, rango_act[1]))
        .order_by('fecha')
    )
    hechos = set(EjercicioHecho.objects.filter(sesion__in=sesiones).values_list('sesion_id', 'ejercicio_id'))
    pesos = {(p.sesion_id, p.ejercicio_id): p for p in PesoEjercicio.objects.filter(sesion__in=sesiones)}

    def ultima(dia, rango):  # sesión más reciente de ese día dentro del rango
        candidatas = [s for s in sesiones if s.dia_id == dia.pk and rango[0] <= s.fecha <= rango[1]]
        return candidatas[-1] if candidatas else None

    filas, carga_ant, carga_act = [], Decimal(0), Decimal(0)
    for d in dias:
        ej = ejercicios[d.pk]
        s_ant, s_act = ultima(d, rango_ant), ultima(d, rango_act)
        ant = _resumen_sesion(s_ant, ej, hechos, pesos) if s_ant else None
        act = _resumen_sesion(s_act, ej, hechos, pesos) if s_act else None
        carga_ant += ant['carga'] if ant else 0
        carga_act += act['carga'] if act else 0
        filas.append({
            'dia': d, 'anterior': ant, 'actual': act,
            'delta': (act['carga'] - ant['carga']) if ant and act else None,
            'ejercicios': [
                {
                    'ejercicio': e,
                    'anterior': pesos.get((s_ant.pk, e.pk)) if s_ant else None,
                    'actual': pesos.get((s_act.pk, e.pk)) if s_act else None,
                }
                for e in ej
                if e.bloque.tipo != Bloque.MOVILIDAD
                and ((s_ant and (s_ant.pk, e.pk) in pesos) or (s_act and (s_act.pk, e.pk) in pesos))
            ],
        })
    return {
        'dias': filas, 'anterior': rango_ant, 'actual': rango_act,
        'carga_anterior': carga_ant, 'carga_actual': carga_act,
        'hay_anterior': any(f['anterior'] for f in filas),
    }


def estado_semana(usuario, rutina, hoy):
    """{dia_id: Sesion} de la semana actual, con la cantidad de ejercicios hechos."""
    lunes = lunes_de(hoy)
    sesiones = Sesion.objects.filter(
        usuario=usuario, dia__rutina=rutina, fecha__range=(lunes, lunes + dt.timedelta(days=6))
    ).annotate(n_hechos=Count('hechos')).order_by('fecha')
    return {s.dia_id: s for s in sesiones}


def serie_peso(usuario):
    """Registros de peso en orden cronológico con variación y media móvil."""
    registros = list(RegistroPeso.objects.filter(usuario=usuario).order_by('fecha'))
    serie = []
    for i, r in enumerate(registros):
        ventana = registros[max(0, i - VENTANA_MEDIA + 1): i + 1]
        media = sum((x.kg for x in ventana), Decimal(0)) / len(ventana)
        serie.append({
            'id': r.pk,
            'fecha': r.fecha,
            'kg': r.kg,
            'nota': r.nota,
            'variacion': (r.kg - registros[i - 1].kg) if i else None,
            'media': media.quantize(Decimal('0.01')),
        })
    return serie


def exportar_pesos_excel(usuario):
    """Genera un .xlsx con el historial de peso y una gráfica; devuelve bytes."""
    serie = serie_peso(usuario)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Peso'

    lila, borde = '7C5CD6', Side(style='thin', color='E4DBFA')
    ws.append(['Fecha', 'Peso (kg)', 'Variación (kg)', f'Media {VENTANA_MEDIA} registros (kg)', 'Nota'])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=lila)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 32

    for fila in serie:
        ws.append([
            fila['fecha'], float(fila['kg']),
            float(fila['variacion']) if fila['variacion'] is not None else None,
            float(fila['media']), fila['nota'],
        ])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = 'DD/MM/YYYY'
        for c in row[1:4]:
            c.number_format = '0.00'
        for c in row:
            c.border = Border(bottom=borde)
            if row[0].row % 2 == 0:
                c.fill = PatternFill('solid', fgColor='F7F3FF')

    for i, ancho in enumerate([14, 12, 15, 22, 40], start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = 'A2'

    if len(serie) >= 2:
        grafica = LineChart()
        grafica.title = 'Evolución del peso'
        grafica.y_axis.title = 'kg'
        grafica.x_axis.number_format = 'dd/mm'
        grafica.height, grafica.width = 8, 18
        ultima = len(serie) + 1
        grafica.add_data(Reference(ws, min_col=2, min_row=1, max_row=ultima), titles_from_data=True)
        grafica.add_data(Reference(ws, min_col=4, min_row=1, max_row=ultima), titles_from_data=True)
        grafica.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ultima))
        grafica.series[0].graphicalProperties.line.solidFill = lila
        grafica.series[1].graphicalProperties.line.solidFill = 'F58FB5'
        ws.add_chart(grafica, 'G2')

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _estilizar_encabezado(ws, ancho_columnas):
    lila = '7C5CD6'
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=lila)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30
    for i, ancho in enumerate(ancho_columnas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = 'A2'


def exportar_registro_excel(usuario, desde=None, hasta=None):
    """Excel con el registro de entrenamientos del usuario entre `desde` y `hasta` (inclusive).

    Hoja "Registro": una fila por ejercicio de cada sesión, con si se hizo y el peso
    (tipo de carga, barra, discos, por mancuerna y total).
    Hoja "Resumen": una fila por sesión con su estado y la carga total.
    """
    sesiones = Sesion.objects.filter(usuario=usuario).select_related('dia').order_by('fecha', 'dia__orden')
    if desde:
        sesiones = sesiones.filter(fecha__gte=desde)
    if hasta:
        sesiones = sesiones.filter(fecha__lte=hasta)
    sesiones = list(sesiones)

    hechos = set(EjercicioHecho.objects.filter(sesion__in=sesiones).values_list('sesion_id', 'ejercicio_id'))
    pesos = {(p.sesion_id, p.ejercicio_id): p for p in PesoEjercicio.objects.filter(sesion__in=sesiones)}
    ejercicios_por_dia = {}
    for s in sesiones:
        if s.dia_id not in ejercicios_por_dia:
            ejercicios_por_dia[s.dia_id] = list(
                Ejercicio.objects.filter(bloque__dia_id=s.dia_id)
                .select_related('bloque').order_by('bloque__orden', 'orden')
            )

    wb = Workbook()
    registro = wb.active
    registro.title = 'Registro'
    registro.append([
        'Fecha', 'Día', 'Enfoque', 'Bloque', 'Ejercicio', 'Series y reps', 'Hecho',
        'Tipo de carga', 'Barra (kg)', 'Discos (kg)', 'Por mancuerna (kg)', 'Total (kg)', 'Indicaciones',
    ])
    resumen = wb.create_sheet('Resumen')
    resumen.append(['Fecha', 'Día', 'Enfoque', 'Estado', 'Ejercicios hechos', 'Total de ejercicios', 'Carga total (kg)'])

    nombres_bloque = dict(Bloque.TIPOS)
    nombres_carga = dict(TIPOS_CARGA)
    numero = lambda v: float(v) if v is not None else None  # noqa: E731
    for s in sesiones:
        ejercicios = ejercicios_por_dia[s.dia_id]
        n_hechos, carga_total = 0, Decimal(0)
        for e in ejercicios:
            hecho = (s.pk, e.pk) in hechos
            n_hechos += hecho
            p = pesos.get((s.pk, e.pk))
            carga_total += p.kg if p else 0
            registro.append([
                s.fecha, s.dia.titulo, s.dia.enfoque, nombres_bloque[e.bloque.tipo], e.nombre,
                e.prescripcion, 'Sí' if hecho else 'No',
                nombres_carga[p.modo] if p else None,
                numero(p.kg_barra) if p else None, numero(p.kg_discos) if p else None,
                numero(p.kg_mancuerna) if p else None, numero(p.kg) if p else None, e.nota,
            ])
        estado = 'Completado' if s.completada else ('En curso' if n_hechos else 'Pendiente')
        resumen.append([s.fecha, s.dia.titulo, s.dia.enfoque, estado, n_hechos, len(ejercicios), float(carga_total)])

    _estilizar_encabezado(registro, [13, 10, 20, 18, 42, 16, 9, 15, 11, 11, 18, 11, 44])
    _estilizar_encabezado(resumen, [13, 10, 20, 14, 18, 20, 16])
    for fila in registro.iter_rows(min_row=2):
        fila[0].number_format = 'DD/MM/YYYY'
        for celda in fila[8:12]:
            celda.number_format = '0.0#'
        fila[6].alignment = Alignment(horizontal='center')
        if fila[6].value == 'Sí':
            fila[6].fill = PatternFill('solid', fgColor='CDF3E4')
    for fila in resumen.iter_rows(min_row=2):
        fila[0].number_format = 'DD/MM/YYYY'
        fila[6].number_format = '0.0#'
    registro.auto_filter.ref = registro.dimensions

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
