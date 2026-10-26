import datetime as dt
import io
from decimal import Decimal as D

import pytest
from django.contrib.auth import get_user_model
from django.core import mail
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import connection
from django.db.migrations.executor import MigrationExecutor
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apptrainer import recordatorios, services
from apptrainer.cargas import armar_carga, lineas_detalle
from apptrainer.models import Dia, Ejercicio, PesoEjercicio, Recordatorio, Sesion

from .test_dia_pesos import dia1, ej  # noqa: F401  (fixtures y helpers)

LUNES = dt.date(2026, 10, 26)           # un lunes: toca el Día 1
AHORA = timezone.make_aware(dt.datetime(2026, 10, 26, 21, 30))


def nombre_en(con_rutina, nombre):
    return Ejercicio.objects.get(bloque__dia__rutina__usuario=con_rutina, nombre=nombre)


def pagina(cliente, dia, **params):
    return cliente.get(reverse('apptrainer:dia', args=[dia.pk]), params).content.decode()


def guardar(cliente, dia, fecha=None, **campos):
    datos = {'fecha': (fecha or timezone.localdate()).isoformat(), **campos}
    return cliente.post(reverse('apptrainer:dia_guardar', args=[dia.pk]), datos)


# ---------- configuración de cada ejercicio, según la lista de Grecia ----------

SOLO_TOTAL = ['Máquina de aducción', 'Drop serie en máquina de abducción', 'Plancha sobre fitball',
              'Elevaciones de cadera en colchoneta', 'Plancha constructora', 'Oblicuos con banda',
              'Flexiones de brazo', 'Fondos en cajón', 'Plancha dinámica con banda en rodillas',
              'Bicho muerto', 'Puente con fitball en rodillas', 'Posteriores en TRX']


@pytest.mark.parametrize('nombre', SOLO_TOTAL)
def test_ejercicios_solo_con_peso_total(con_rutina, nombre):
    e = nombre_en(con_rutina, nombre)
    assert e.modos == ['total'] and e.carga_habitual == 'total'


@pytest.mark.parametrize('nombre, modos, habitual, por_lado', [
    ('Hip thrust', ['barra'], 'barra', True),
    ('Remo horizontal en máquina', ['lados'], 'lados', False),
    ('Sentadilla búlgara', ['mancuernas'], 'mancuernas', False),
    ('Step up', ['mancuernas', 'total'], 'mancuernas', False),
])
def test_ejercicios_con_carga_especifica(con_rutina, nombre, modos, habitual, por_lado):
    e = nombre_en(con_rutina, nombre)
    assert (e.modos, e.carga_habitual, e.carga_por_lado) == (modos, habitual, por_lado)


@pytest.mark.parametrize('nombre', ['Empuje de rodillas en landmine', 'Bíceps martillo iso alternado',
                                    'Peso muerto rumano', 'Máquina de espalda', 'Vuelos laterales'])
def test_los_demas_ejercicios_no_cambian(con_rutina, nombre):
    assert nombre_en(con_rutina, nombre).modos == ['barra', 'mancuernas', 'total']


@pytest.mark.parametrize('nombre', SOLO_TOTAL)
def test_en_pantalla_solo_total_no_hay_campos_de_barra_ni_mancuerna(cliente, con_rutina, nombre):
    e = nombre_en(con_rutina, nombre)
    html = pagina(cliente, e.bloque.dia)
    assert f'name="kg_{e.pk}"' in html
    for campo in ('barra', 'discos', 'lado', 'sololado', 'mancuerna'):
        assert f'name="{campo}_{e.pk}"' not in html
    assert f'id="modo_{e.pk}_barra"' not in html and f'id="modo_{e.pk}_mancuernas"' not in html


def test_en_pantalla_el_selector_solo_aparece_si_hay_opciones(cliente, con_rutina):
    html = pagina(cliente, Dia.objects.get(rutina__usuario=con_rutina, orden=3))
    step = nombre_en(con_rutina, 'Step up')
    bulgara = nombre_en(con_rutina, 'Sentadilla búlgara')
    assert f'id="modo_{step.pk}_mancuernas"' in html and f'id="modo_{step.pk}_total"' in html
    assert f'id="modo_{step.pk}_barra"' not in html
    assert f'name="mancuerna_{bulgara.pk}"' in html and f'name="kg_{bulgara.pk}"' not in html


def test_el_servidor_ignora_tipos_de_carga_que_el_ejercicio_no_ofrece(cliente, dia1):
    aduccion = ej(dia1, 'Máquina de aducción')
    guardar(cliente, dia1, **{f'modo_{aduccion.pk}': 'mancuernas', f'mancuerna_{aduccion.pk}': '10',
                              f'barra_{aduccion.pk}': '20', f'kg_{aduccion.pk}': '35'})
    p = PesoEjercicio.objects.get()
    assert (p.modo, p.kg, p.kg_mancuerna, p.kg_barra) == ('total', D(35), None, None)


# ---------- remo: discos de cada lado (nuevo tipo de carga) ----------

def test_discos_por_lado_cuentan_dos_veces():
    c = armar_carga('lados', por_lado=D(20))
    assert (c.modo, c.kg, c.por_lado, c.barra) == ('lados', 40, 20, None)
    assert armar_carga('lados') is None


def test_guardar_remo_con_discos_de_cada_lado(cliente, con_rutina):
    remo = nombre_en(con_rutina, 'Remo horizontal en máquina')
    dia2 = remo.bloque.dia
    html = pagina(cliente, dia2)
    assert f'name="sololado_{remo.pk}"' in html and 'Discos de cada lado (kg)' in html
    r = guardar(cliente, dia2, **{f'modo_{remo.pk}': 'lados', f'sololado_{remo.pk}': '17,5'})
    assert r.status_code == 302
    p = PesoEjercicio.objects.get()
    assert (p.modo, p.kg, p.kg_por_lado, p.kg_barra) == ('lados', D(35), D('17.5'), None)
    assert p.detalle == '2 × 17,5 = 35 kg'
    assert 'value="17,5"' in pagina(cliente, dia2) and 'Total: <b data-total>35</b>' in pagina(cliente, dia2)


def test_remo_valida_y_conserva_lo_tipeado(cliente, con_rutina):
    remo = nombre_en(con_rutina, 'Remo horizontal en máquina')
    r = guardar(cliente, remo.bloque.dia, **{f'sololado_{remo.pk}': 'abc'})
    assert r.status_code == 422 and not PesoEjercicio.objects.exists() and 'value="abc"' in r.content.decode()
    r = guardar(cliente, remo.bloque.dia, **{f'sololado_{remo.pk}': '300'})
    assert r.status_code == 422 and 'no puede superar 500' in r.content.decode()


def test_excel_incluye_discos_por_lado(cliente, con_rutina):
    remo = nombre_en(con_rutina, 'Remo horizontal en máquina')
    guardar(cliente, remo.bloque.dia, **{f'sololado_{remo.pk}': '20'})
    wb = load_workbook(io.BytesIO(cliente.get(reverse('apptrainer:semana_excel')).content))
    fila = next(f for f in wb['Registro'].iter_rows(min_row=2, values_only=True) if f[4] == 'Remo horizontal en máquina')
    assert fila[7:13] == ('Discos por lado', None, None, 20.0, None, 40.0)


# ---------- "Anterior": cada dato por separado ----------

def test_lineas_de_detalle_por_tipo():
    assert lineas_detalle('barra', D(20), None, None, D(50), D(15)) == [
        ('Barra', '20 kg'), ('Disco de cada lado', '15 kg'), ('Total (barra + 2 lados)', '50 kg')]
    assert lineas_detalle('barra', D(20), D(30), None, D(50)) == [
        ('Barra', '20 kg'), ('Discos (los dos lados)', '30 kg'), ('Total', '50 kg')]
    assert lineas_detalle('mancuernas', None, None, D('12.5'), D(25)) == [
        ('Cada mancuerna', '12,5 kg'), ('Total (2 mancuernas)', '25 kg')]
    assert lineas_detalle('lados', None, None, None, D(40), D(20)) == [
        ('Discos de cada lado', '20 kg'), ('Total (2 lados)', '40 kg')]
    assert lineas_detalle('total', None, None, None, D(35)) == [('Peso total', '35 kg')]


def test_anterior_detalla_barra_y_disco_de_cada_lado(cliente, dia1, usuario):
    hip = ej(dia1, 'Hip thrust')
    antes = timezone.localdate() - dt.timedelta(days=4)
    services.guardar_pesos(usuario, dia1, antes, {hip.pk: armar_carga('barra', barra=D(20), por_lado=D(15), lados=True)})
    html = pagina(cliente, dia1)
    assert 'Anterior · ' in html
    for linea in ('<dt>Barra</dt><dd>20 kg</dd>', '<dt>Disco de cada lado</dt><dd>15 kg</dd>',
                  '<dt>Total (barra + 2 lados)</dt><dd>50 kg</dd>'):
        assert linea in html


def test_anterior_detalla_mancuernas_y_discos_por_lado(cliente, con_rutina, usuario):
    antes = timezone.localdate() - dt.timedelta(days=4)
    vuelos = nombre_en(con_rutina, 'Vuelos laterales')
    remo = nombre_en(con_rutina, 'Remo horizontal en máquina')
    services.guardar_pesos(usuario, vuelos.bloque.dia, antes, {vuelos.pk: armar_carga('mancuernas', mancuerna=D(4))})
    services.guardar_pesos(usuario, remo.bloque.dia, antes, {remo.pk: armar_carga('lados', por_lado=D(20))})
    assert '<dt>Cada mancuerna</dt><dd>4 kg</dd>' in pagina(cliente, vuelos.bloque.dia)
    assert '<dt>Discos de cada lado</dt><dd>20 kg</dd>' in pagina(cliente, remo.bloque.dia)


# ---------- notas de la semana pasada, destacadas ----------

def _nota(usuario, dia, fecha, texto):
    services.guardar_pesos(usuario, dia, fecha, {}, comentario=texto)


def test_la_pantalla_principal_destaca_las_notas_de_la_semana_pasada(cliente, con_rutina, usuario):
    hoy = timezone.localdate()
    lunes = services.lunes_de(hoy)
    d1, d2 = (Dia.objects.get(rutina__usuario=con_rutina, orden=n) for n in (1, 2))
    Dia.objects.filter(rutina__usuario=con_rutina).update(dia_semana=None)
    Dia.objects.filter(pk=d2.pk).update(dia_semana=hoy.weekday())  # hoy toca el Día 2
    _nota(usuario, d1, lunes - dt.timedelta(days=7), 'Día 1: subir peso en sumo')
    _nota(usuario, d2, lunes - dt.timedelta(days=5), 'Día 2: dolor en el hombro')
    _nota(usuario, d1, lunes - dt.timedelta(days=15), 'Nota vieja, de hace dos semanas')
    html = cliente.get(reverse('apptrainer:hoy')).content.decode()
    assert 'Tus notas de la semana pasada' in html
    assert html.index('Tus notas de la semana pasada') < html.index('class="card hero"')   # arriba de todo
    assert 'Día 1: subir peso en sumo' in html and 'Día 2: dolor en el hombro' in html
    assert 'Nota vieja' not in html
    assert html.index('Día 2: dolor en el hombro') < html.index('Día 1: subir peso en sumo')  # primero el de hoy
    assert 'Hoy toca' in html and 'nota-hoy' in html


def test_sin_notas_no_hay_tarjeta_destacada(cliente, con_rutina):
    assert 'Tus notas de la semana pasada' not in cliente.get(reverse('apptrainer:hoy')).content.decode()


def test_se_muestra_la_nota_mas_reciente_de_cada_dia(usuario, dia1):
    lunes = services.lunes_de(timezone.localdate())
    _nota(usuario, dia1, lunes - dt.timedelta(days=6), 'primera')
    _nota(usuario, dia1, lunes - dt.timedelta(days=2), 'segunda')
    notas = services.comentarios_semana_pasada(usuario, dia1.rutina, timezone.localdate())
    assert [n['comentario'] for n in notas] == ['segunda']


def test_la_pantalla_del_dia_destaca_la_nota_de_la_semana_anterior(cliente, dia1, usuario):
    hoy = timezone.localdate()
    _nota(usuario, dia1, services.lunes_de(hoy) - dt.timedelta(days=4), 'Cuidar la rodilla derecha')
    html = pagina(cliente, dia1)
    assert 'Tu nota de la semana pasada' in html and 'Cuidar la rodilla derecha' in html
    assert html.index('Tu nota de la semana pasada') < html.index('class="date-bar"')
    # mirando otra semana, la "semana anterior" es la de esa fecha
    assert 'Cuidar la rodilla derecha' not in pagina(cliente, dia1, fecha=(hoy + dt.timedelta(days=14)).isoformat())


def test_las_notas_son_privadas(client, usuario, otro_usuario, dia1):
    _nota(usuario, dia1, services.lunes_de(timezone.localdate()) - dt.timedelta(days=3), 'secreto')
    call_command('seed_rutina', usuario='otra', stdout=io.StringIO())
    client.force_login(otro_usuario)
    assert 'secreto' not in client.get(reverse('apptrainer:hoy')).content.decode()


# ---------- recordatorios por correo ----------

@pytest.fixture
def con_correo(con_rutina, usuario, settings):
    usuario.email = 'grecia@example.com'
    usuario.save()
    settings.RECORDATORIO_HORA = 21
    settings.EMAIL_HOST = 'smtp.ejemplo.test'   # el envío real lo intercepta el buzón de pruebas
    settings.APP_URL = 'https://apptrainer.onrender.com'
    mail.outbox.clear()
    return usuario


def _hecho(usuario, nombre, fecha=LUNES):
    e = Ejercicio.objects.filter(bloque__dia__rutina__usuario=usuario, nombre=nombre).order_by('bloque__dia__orden').first()
    services.alternar_ejercicio(usuario, e, fecha)
    return e


def test_avisa_si_terminaste_el_dia_y_faltan_pesos(con_correo):
    _hecho(con_correo, 'Hip thrust')
    _hecho(con_correo, 'Peso muerto rumano')
    e = _hecho(con_correo, 'Sentadilla goblet')
    services.guardar_pesos(con_correo, e.bloque.dia, LUNES, {e.pk: armar_carga('total', total=D(12))})
    resultados = recordatorios.enviar_pendientes(AHORA)
    assert [r['estado'] for r in resultados] == ['enviado']
    assert [x.nombre for x in resultados[0]['pendientes']] == ['Hip thrust', 'Peso muerto rumano']
    (correo,) = mail.outbox
    assert correo.to == ['grecia@example.com'] and 'Día 1' in correo.subject
    assert 'Hip thrust' in correo.body and 'Peso muerto rumano' in correo.body and 'Sentadilla goblet' not in correo.body
    assert 'https://apptrainer.onrender.com/app/dia/' in correo.body and 'fecha=2026-10-26' in correo.body
    assert Recordatorio.objects.get().fecha == LUNES


def test_el_movilidad_no_cuenta_como_peso_pendiente(con_correo):
    _hecho(con_correo, 'Movilidad de tobillo')
    assert recordatorios.enviar_pendientes(AHORA) == [] and not mail.outbox


def test_no_avisa_antes_de_la_hora(con_correo):
    _hecho(con_correo, 'Hip thrust')
    temprano = timezone.make_aware(dt.datetime(2026, 10, 26, 18, 0))
    assert recordatorios.enviar_pendientes(temprano) == [] and not mail.outbox
    assert len(recordatorios.enviar_pendientes(temprano, ignorar_hora=True)) == 1


def test_no_avisa_dos_veces_el_mismo_dia(con_correo):
    _hecho(con_correo, 'Hip thrust')
    recordatorios.enviar_pendientes(AHORA)
    assert recordatorios.enviar_pendientes(AHORA + dt.timedelta(hours=1)) == []
    assert len(mail.outbox) == 1


def test_no_avisa_si_ya_anotaste_todos_los_pesos(con_correo):
    e = _hecho(con_correo, 'Hip thrust')
    services.guardar_pesos(con_correo, e.bloque.dia, LUNES, {e.pk: armar_carga('barra', barra=D(20), por_lado=D(10), lados=True)})
    assert recordatorios.enviar_pendientes(AHORA) == [] and not mail.outbox


def test_no_avisa_si_no_hubo_actividad(con_correo):
    assert recordatorios.enviar_pendientes(AHORA) == [] and not mail.outbox


def test_no_avisa_en_dias_sin_entrenamiento(con_correo):
    _hecho(con_correo, 'Hip thrust')                     # el Día 1 es de los lunes
    martes = timezone.make_aware(dt.datetime(2026, 10, 27, 21, 30))
    assert recordatorios.enviar_pendientes(martes) == []


def test_no_avisa_sin_correo_ni_a_cuentas_inactivas(con_correo):
    _hecho(con_correo, 'Hip thrust')
    con_correo.email = ''
    con_correo.save()
    assert recordatorios.enviar_pendientes(AHORA) == []
    con_correo.email, con_correo.is_active = 'grecia@example.com', False
    con_correo.save()
    assert recordatorios.enviar_pendientes(AHORA) == [] and not mail.outbox


def test_simular_no_envia_ni_registra(con_correo):
    _hecho(con_correo, 'Hip thrust')
    assert recordatorios.enviar_pendientes(AHORA, simular=True)[0]['estado'] == 'simulado'
    assert not mail.outbox and not Recordatorio.objects.exists()


def test_un_error_de_correo_no_registra_el_aviso_y_se_reintenta(con_correo, monkeypatch):
    _hecho(con_correo, 'Hip thrust')
    monkeypatch.setattr(recordatorios, 'send_mail', lambda *a, **k: (_ for _ in ()).throw(OSError('sin red')))
    assert recordatorios.enviar_pendientes(AHORA)[0]['estado'] == 'error'
    assert not Recordatorio.objects.exists()
    monkeypatch.undo()
    assert recordatorios.enviar_pendientes(AHORA + dt.timedelta(hours=1))[0]['estado'] == 'enviado'


def test_un_fallo_con_un_usuario_no_frena_a_los_demas(con_correo, otro_usuario, monkeypatch):
    call_command('seed_rutina', usuario='otra', stdout=io.StringIO())
    otro_usuario.email = 'otra@example.com'
    otro_usuario.save()
    _hecho(con_correo, 'Hip thrust')
    _hecho(otro_usuario, 'Hip thrust')
    real = recordatorios.send_mail
    monkeypatch.setattr(recordatorios, 'send_mail', lambda asunto, cuerpo, de, para, **k: (
        (_ for _ in ()).throw(OSError('x')) if para == ['grecia@example.com'] else real(asunto, cuerpo, de, para, **k)))
    estados = {r['usuario'].username: r['estado'] for r in recordatorios.enviar_pendientes(AHORA)}
    assert estados == {'grecia': 'error', 'otra': 'enviado'}
    assert [m.to for m in mail.outbox] == [['otra@example.com']]


def test_cada_usuario_recibe_solo_lo_suyo(con_correo, otro_usuario):
    call_command('seed_rutina', usuario='otra', stdout=io.StringIO())
    otro_usuario.email = 'otra@example.com'
    otro_usuario.save()
    _hecho(otro_usuario, 'Peso muerto rumano')
    recordatorios.enviar_pendientes(AHORA)
    (correo,) = mail.outbox
    assert correo.to == ['otra@example.com'] and 'Peso muerto rumano' in correo.body and 'Hip thrust' not in correo.body


def test_comando_simular(con_correo):
    hoy = timezone.localdate()
    Dia.objects.filter(rutina__usuario=con_correo).update(dia_semana=None)
    Dia.objects.filter(rutina__usuario=con_correo, orden=1).update(dia_semana=hoy.weekday())
    _hecho(con_correo, 'Hip thrust', hoy)
    salida = io.StringIO()
    call_command('enviar_recordatorios', '--simular', '--ahora', stdout=salida)
    assert 'SIMULADO: grecia · Día 1 · faltan: Hip thrust' in salida.getvalue() and not mail.outbox


def test_comando_envia_y_luego_no_repite(con_correo):
    hoy = timezone.localdate()
    Dia.objects.filter(rutina__usuario=con_correo).update(dia_semana=None)
    Dia.objects.filter(rutina__usuario=con_correo, orden=1).update(dia_semana=hoy.weekday())
    _hecho(con_correo, 'Hip thrust', hoy)
    salida = io.StringIO()
    call_command('enviar_recordatorios', '--ahora', stdout=salida)
    call_command('enviar_recordatorios', '--ahora', stdout=salida)
    assert 'ENVIADO: grecia' in salida.getvalue() and 'Sin recordatorios para enviar.' in salida.getvalue()
    assert len(mail.outbox) == 1


def test_comando_exige_correo_configurado_en_produccion(con_correo, settings):
    settings.DEBUG, settings.EMAIL_HOST = False, ''
    with pytest.raises(CommandError, match='EMAIL_HOST'):
        call_command('enviar_recordatorios')
    call_command('enviar_recordatorios', '--simular', stdout=io.StringIO())   # simular siempre se puede


def test_comando_con_usuario_inexistente(con_correo):
    with pytest.raises(CommandError, match='No existe'):
        call_command('enviar_recordatorios', '--usuario', 'nadie')


# ---------- el historial sobrevive a estas migraciones ----------

@pytest.mark.django_db(transaction=True)
def test_migraciones_de_octubre_conservan_el_historial_y_reconfiguran_los_ejercicios():
    antes = [('apptrainer', '0007_hip_thrust_barra_y_cada_lado')]
    MigrationExecutor(connection).migrate(antes)
    try:
        apps = MigrationExecutor(connection).loader.project_state(antes).apps
        User, Rutina, Dia_, Bloque = (apps.get_model('auth', 'User'), apps.get_model('apptrainer', 'Rutina'),
                                      apps.get_model('apptrainer', 'Dia'), apps.get_model('apptrainer', 'Bloque'))
        Ejercicio_, Sesion_, Hecho, Peso = (apps.get_model('apptrainer', n) for n in
                                            ('Ejercicio', 'Sesion', 'EjercicioHecho', 'PesoEjercicio'))
        u = User.objects.create(username='grecia')
        rutina = Rutina.objects.create(usuario=u, nombre='Rutina semanal', activa=True)
        dia = Dia_.objects.create(rutina=rutina, orden=1, titulo='Día 1')
        bloque = Bloque.objects.create(dia=dia, tipo='principal', orden=1)
        aduccion = Ejercicio_.objects.create(bloque=bloque, orden=1, nombre='Máquina de aducción', imagen='aduccion')
        remo = Ejercicio_.objects.create(bloque=bloque, orden=2, nombre='Remo horizontal en máquina', imagen='remo-maquina')
        hip = Ejercicio_.objects.create(bloque=bloque, orden=3, nombre='Hip thrust', nota='Barra normal',
                                        modos_permitidos='barra,total', carga_por_lado=True)
        landmine = Ejercicio_.objects.create(bloque=bloque, orden=4, nombre='Empuje de rodillas en landmine')
        s = Sesion_.objects.create(usuario=u, dia=dia, fecha='2026-10-19', completada=True, comentario='buena sesión')
        for e in (aduccion, remo, hip, landmine):
            Hecho.objects.create(sesion=s, ejercicio=e)
        # lo que se cargó con las opciones viejas queda tal cual
        Peso.objects.create(sesion=s, ejercicio=aduccion, modo='barra', kg_barra=20, kg_discos=10, kg=30)
        Peso.objects.create(sesion=s, ejercicio=remo, modo='mancuernas', kg_mancuerna=15, kg=30)
        Peso.objects.create(sesion=s, ejercicio=hip, modo='barra', kg_barra=20, kg_por_lado=15, kg=50)
        Peso.objects.create(sesion=s, ejercicio=landmine, modo='barra', kg_barra=20, kg_discos=10, kg=30)
        foto = sorted(Peso.objects.values_list('ejercicio__nombre', 'modo', 'kg_barra', 'kg_discos', 'kg_por_lado', 'kg_mancuerna', 'kg'))

        ultimo = MigrationExecutor(connection)
        ultimo.migrate(ultimo.loader.graph.leaf_nodes())

        from apptrainer import models as m
        assert sorted(m.PesoEjercicio.objects.values_list('ejercicio__nombre', 'modo', 'kg_barra', 'kg_discos', 'kg_por_lado', 'kg_mancuerna', 'kg')) == foto
        sesion = m.Sesion.objects.get()
        assert (sesion.comentario, sesion.completada, str(sesion.fecha)) == ('buena sesión', True, '2026-10-19')
        assert m.EjercicioHecho.objects.count() == 4
        # los ejercicios se reconfiguran según la lista
        cfg = {e.nombre: (e.modos, e.carga_habitual) for e in m.Ejercicio.objects.all()}
        assert cfg['Máquina de aducción'] == (['total'], 'total')
        assert cfg['Remo horizontal en máquina'] == (['lados'], 'lados')
        assert cfg['Hip thrust'] == (['barra'], 'barra')
        assert cfg['Empuje de rodillas en landmine'] == (['barra', 'mancuernas', 'total'], 'total')  # sin cambios
        assert not hasattr(m.Ejercicio, 'imagen') and m.Recordatorio.objects.count() == 0
        # lo viejo se sigue mostrando con su detalle original
        assert [p.detalle for p in m.PesoEjercicio.objects.order_by('pk')] == [
            '20 + 10 = 30 kg', '2 × 15 = 30 kg', '20 + 2 × 15 = 50 kg', '20 + 10 = 30 kg']
    finally:
        final = MigrationExecutor(connection)
        final.migrate(final.loader.graph.leaf_nodes())
