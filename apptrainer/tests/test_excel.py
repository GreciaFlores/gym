import datetime as dt
import io
from decimal import Decimal

from django.urls import reverse
from openpyxl import load_workbook

from apptrainer import services
from apptrainer.models import RegistroPeso


def _crear(usuario, pares):
    base = dt.date(2026, 1, 5)
    for i, kg in enumerate(pares):
        RegistroPeso.objects.create(usuario=usuario, fecha=base + dt.timedelta(days=7 * i), kg=Decimal(kg))


def test_serie_calcula_variacion_y_media_movil(usuario):
    _crear(usuario, ['70', '69', '68', '67', '66'])
    serie = services.serie_peso(usuario)
    assert serie[0]['variacion'] is None
    assert serie[1]['variacion'] == Decimal('-1')
    assert serie[1]['media'] == Decimal('69.50')            # (70+69)/2
    assert serie[4]['media'] == Decimal('67.50')            # (69+68+67+66)/4: ventana de 4


def test_excel_tiene_encabezados_filas_y_grafica(usuario):
    _crear(usuario, ['70', '69.5', '69'])
    wb = load_workbook(io.BytesIO(services.exportar_pesos_excel(usuario)))
    ws = wb['Peso']
    assert [c.value for c in ws[1]][:3] == ['Fecha', 'Peso (kg)', 'Variación (kg)']
    assert ws.max_row == 4
    assert ws['B2'].value == 70.0 and ws['C2'].value is None
    assert ws['C3'].value == -0.5
    assert ws['A2'].value.date() == dt.date(2026, 1, 5)
    assert len(ws._charts) == 1


def test_excel_sin_registros_no_falla(usuario):
    wb = load_workbook(io.BytesIO(services.exportar_pesos_excel(usuario)))
    assert wb['Peso'].max_row == 1


def test_descarga_excel_por_http(cliente, usuario):
    _crear(usuario, ['70', '69'])
    r = cliente.get(reverse('apptrainer:peso_excel'))
    assert r.status_code == 200
    assert r['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'attachment; filename="apptrainer-peso-' in r['Content-Disposition']
    assert load_workbook(io.BytesIO(r.content))['Peso'].max_row == 3


def test_excel_solo_incluye_pesos_propios(cliente, usuario, otro_usuario):
    _crear(usuario, ['70'])
    _crear(otro_usuario, ['50', '51'])
    wb = load_workbook(io.BytesIO(cliente.get(reverse('apptrainer:peso_excel')).content))
    assert wb['Peso'].max_row == 2
