import datetime as dt
import io
from decimal import Decimal as D

import pytest
from django.core.management import call_command
from django.core.management.base import CommandError
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apptrainer import services
from apptrainer.cargas import armar_carga, detalle_texto, kg_texto
from apptrainer.catalogo import CATALOGO, DESCRIPCIONES
from apptrainer.models import Dia, Ejercicio, PesoEjercicio, Sesion

from .test_dia_pesos import dia1, ej  # noqa: F401  (fixtures y helpers)


def guardar_modo(cliente, dia, fecha=None, **campos):
    datos = {'fecha': (fecha or timezone.localdate()).isoformat()}
    datos.update(campos)
    return cliente.post(reverse('apptrainer:dia_guardar', args=[dia.pk]), datos)


# ---------- cálculo de la carga ----------

def test_barra_mas_discos():
    c = armar_carga('barra', barra=D(20), discos=D(30))
    assert c.kg == 50 and detalle_texto('barra', D(20), D(30), None, D(50)) == '20 + 30 = 50 kg'


def test_barra_sola_o_solo_discos_alcanza():
    assert armar_carga('barra', barra=D(20)).kg == 20
    assert armar_carga('barra', discos=D(15)).kg == 15


def test_mancuernas_cuentan_dos():
    c = armar_carga('mancuernas', mancuerna=D('12.5'))
    assert c.kg == 25 and detalle_texto('mancuernas', None, None, D('12.5'), D(25)) == '2 × 12,5 = 25 kg'


@pytest.mark.parametrize('modo', ['barra', 'mancuernas', 'total'])
def test_sin_datos_no_hay_carga(modo):
    assert armar_carga(modo) is None


def test_total_maximo():
    with pytest.raises(ValueError, match='500'):
        armar_carga('mancuernas', mancuerna=D(300))
    assert armar_carga('total', total=D(500)).kg == 500


def test_kg_texto_sin_ceros_de_mas():
    assert [kg_texto(D(x)) for x in ('52.00', '52.50', '20', '0.25')] == ['52', '52,5', '20', '0,25']


# ---------- guardar por tipo de carga ----------

def test_guardar_barra_con_discos(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    r = guardar_modo(cliente, dia1, **{f'modo_{hip.pk}': 'barra', f'barra_{hip.pk}': '20', f'discos_{hip.pk}': '30'})
    assert r.status_code == 302
    p = PesoEjercicio.objects.get()
    assert (p.modo, p.kg, p.kg_barra, p.kg_discos, p.kg_mancuerna) == ('barra', D(50), D(20), D(30), None)
    assert p.detalle == '20 + 30 = 50 kg'


def test_guardar_mancuernas(cliente, dia1):
    vuelos = ej(dia1.rutina.dias.get(orden=3), 'Vuelos laterales')
    guardar_modo(cliente, vuelos.bloque.dia, **{f'modo_{vuelos.pk}': 'mancuernas', f'mancuerna_{vuelos.pk}': '4,5'})
    p = PesoEjercicio.objects.get()
    assert (p.modo, p.kg, p.kg_mancuerna) == ('mancuernas', D(9), D('4.5'))
    assert p.detalle == '2 × 4,5 = 9 kg'


def test_cambiar_de_modo_actualiza_el_registro(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    guardar_modo(cliente, dia1, **{f'modo_{hip.pk}': 'barra', f'barra_{hip.pk}': '20', f'discos_{hip.pk}': '30'})
    guardar_modo(cliente, dia1, **{f'modo_{hip.pk}': 'total', f'kg_{hip.pk}': '60'})
    p = PesoEjercicio.objects.get()
    assert (p.modo, p.kg, p.kg_barra, p.kg_discos) == ('total', D(60), None, None)


def test_solo_se_validan_los_campos_del_modo_elegido(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    # basura en "mancuerna" no importa si el modo es barra
    guardar_modo(cliente, dia1, **{f'modo_{hip.pk}': 'barra', f'barra_{hip.pk}': '20', f'mancuerna_{hip.pk}': 'xx'})
    assert PesoEjercicio.objects.get().kg == 20


def test_error_en_un_campo_conserva_todo_lo_tipeado(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    r = guardar_modo(cliente, dia1, **{f'modo_{hip.pk}': 'barra', f'barra_{hip.pk}': '20', f'discos_{hip.pk}': 'abc'})
    html = r.content.decode()
    assert r.status_code == 422 and not PesoEjercicio.objects.exists()
    assert 'value="20"' in html and 'value="abc"' in html and 'field-error' in html


def test_total_de_mancuernas_no_puede_pasar_el_maximo(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    r = guardar_modo(cliente, dia1, **{f'modo_{hip.pk}': 'mancuernas', f'mancuerna_{hip.pk}': '300'})
    assert r.status_code == 422 and 'no puede superar 500' in r.content.decode()


def test_modo_invalido_se_trata_como_total(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    guardar_modo(cliente, dia1, **{f'modo_{hip.pk}': 'hackeado', f'kg_{hip.pk}': '40'})
    assert PesoEjercicio.objects.get().modo == 'total'


def test_el_modo_preseleccionado_es_el_habitual_del_ejercicio(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')  # barra por catálogo
    html = cliente.get(reverse('apptrainer:dia', args=[dia1.pk])).content.decode()
    assert f'id="modo_{hip.pk}_barra" name="modo_{hip.pk}" value="barra" form="form-guardar" checked' in html
    rumano = ej(dia1, 'Plancha sobre fitball')  # total
    assert f'id="modo_{rumano.pk}_total" name="modo_{rumano.pk}" value="total" form="form-guardar" checked' in html


def test_el_modo_de_la_vez_anterior_manda_sobre_el_habitual(cliente, dia1):
    plancha = ej(dia1, 'Plancha sobre fitball')
    ayer = timezone.localdate() - dt.timedelta(days=2)
    guardar_modo(cliente, dia1, fecha=ayer, **{f'modo_{plancha.pk}': 'mancuernas', f'mancuerna_{plancha.pk}': '5'})
    html = cliente.get(reverse('apptrainer:dia', args=[dia1.pk])).content.decode()
    assert f'id="modo_{plancha.pk}_mancuernas" name="modo_{plancha.pk}" value="mancuernas" form="form-guardar" checked' in html
    assert 'Anterior: 2 × 5 = 10 kg' in html


def test_resumen_muestra_el_detalle(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    guardar_modo(cliente, dia1, **{f'modo_{hip.pk}': 'barra', f'barra_{hip.pk}': '20', f'discos_{hip.pk}': '30'})
    html = cliente.get(reverse('apptrainer:dia_resumen', args=[dia1.pk])).content.decode()
    assert '20 + 30 = 50 kg' in html and '50 kg' in html


# ---------- pantalla principal: cierre semanal ----------

def _sesion_con_peso(usuario, dia, fecha, ejercicio, modo='total', **kw):
    carga = armar_carga(modo, **kw)
    services.guardar_pesos(usuario, dia, fecha, {ejercicio.pk: carga})
    return Sesion.objects.get(usuario=usuario, dia=dia, fecha=fecha)


def test_cierre_compara_semana_pasada_con_esta(cliente, dia1, usuario):
    hoy = timezone.localdate()
    lunes = services.lunes_de(hoy)
    hip, rumano = ej(dia1, 'Hip thrust'), ej(dia1, 'Peso muerto rumano')
    _sesion_con_peso(usuario, dia1, lunes - dt.timedelta(days=5), hip, 'barra', barra=D(20), discos=D(30))
    _sesion_con_peso(usuario, dia1, lunes - dt.timedelta(days=5), rumano, 'total', total=D(40))
    _sesion_con_peso(usuario, dia1, lunes, hip, 'barra', barra=D(20), discos=D(40))

    cierre = services.cierre_semanal(usuario, dia1.rutina, hoy)
    assert cierre['hay_anterior'] is True
    assert cierre['carga_anterior'] == 90 and cierre['carga_actual'] == 60
    fila = cierre['dias'][0]
    assert fila['anterior']['carga'] == 90 and fila['actual']['carga'] == 60 and fila['delta'] == -30
    detalle = {x['ejercicio'].nombre: (x['anterior'], x['actual']) for x in fila['ejercicios']}
    assert detalle['Hip thrust'][0].detalle == '20 + 30 = 50 kg' and detalle['Hip thrust'][1].detalle == '20 + 40 = 60 kg'
    assert detalle['Peso muerto rumano'][1] is None
    assert cierre['dias'][1]['anterior'] is None and cierre['dias'][1]['delta'] is None


def test_cierre_toma_la_ultima_sesion_del_dia_en_la_semana(usuario, dia1):
    hoy = timezone.localdate()
    lunes = services.lunes_de(hoy)
    hip = ej(dia1, 'Hip thrust')
    _sesion_con_peso(usuario, dia1, lunes - dt.timedelta(days=6), hip, 'total', total=D(30))
    _sesion_con_peso(usuario, dia1, lunes - dt.timedelta(days=2), hip, 'total', total=D(35))
    assert services.cierre_semanal(usuario, dia1.rutina, hoy)['dias'][0]['anterior']['carga'] == 35


def test_cierre_ignora_movilidad_y_otras_semanas(usuario, dia1):
    hoy = timezone.localdate()
    hip = ej(dia1, 'Hip thrust')
    _sesion_con_peso(usuario, dia1, services.lunes_de(hoy) - dt.timedelta(days=15), hip, 'total', total=D(99))
    cierre = services.cierre_semanal(usuario, dia1.rutina, hoy)
    assert cierre['hay_anterior'] is False and cierre['carga_anterior'] == 0


def test_cierre_solo_datos_propios(usuario, otro_usuario, dia1):
    hoy = timezone.localdate()
    _sesion_con_peso(usuario, dia1, services.lunes_de(hoy), ej(dia1, 'Hip thrust'), 'total', total=D(50))
    call_command('seed_rutina', usuario='otra', stdout=open('/dev/null', 'w'))
    rutina_ajena = services.rutina_activa(otro_usuario)
    assert services.cierre_semanal(otro_usuario, rutina_ajena, hoy)['carga_actual'] == 0
    assert services.cierre_semanal(usuario, dia1.rutina, hoy)['carga_actual'] == 50


def test_pantalla_principal_muestra_cierre_con_total_y_detalle(cliente, dia1, usuario):
    hoy = timezone.localdate()
    hip = ej(dia1, 'Hip thrust')
    _sesion_con_peso(usuario, dia1, services.lunes_de(hoy) - dt.timedelta(days=4), hip, 'barra', barra=D(20), discos=D(30))
    html = cliente.get(reverse('apptrainer:hoy')).content.decode()
    assert 'Cierre semanal' in html and 'id="vista-total"' in html and 'id="vista-detalle"' in html
    assert '20 + 30 = 50 kg' in html          # vista detalle
    assert 'class="v-total">50 kg' in html    # vista total
    assert 'Hip thrust' in html and 'Sin registro' in html


def test_pantalla_principal_sin_historial_no_falla(cliente, con_rutina):
    html = cliente.get(reverse('apptrainer:hoy')).content.decode()
    assert 'Todavía no hay registro de la semana pasada' in html


def test_el_menu_ya_no_tiene_peso_corporal(cliente, con_rutina):
    for nombre in ('hoy', 'semana'):
        html = cliente.get(reverse(f'apptrainer:{nombre}')).content.decode()
        nav = html[html.index('<nav class="tabbar"'):html.index('</nav>')]
        assert 'Peso' not in nav and reverse('apptrainer:peso') not in nav
    assert 'kg' not in cliente.get(reverse('apptrainer:hoy')).content.decode().split('Cierre semanal')[0]


# ---------- exportar con rango de fechas ----------

def _filas(resp, hoja='Registro'):
    return list(load_workbook(io.BytesIO(resp.content))[hoja].iter_rows(min_row=2, values_only=True))


def test_exportar_por_rango_de_fechas(cliente, dia1, usuario):
    hip = ej(dia1, 'Hip thrust')
    d1, d2, d3 = dt.date(2026, 3, 2), dt.date(2026, 3, 9), dt.date(2026, 3, 16)
    for f, kg in ((d1, 30), (d2, 40), (d3, 50)):
        _sesion_con_peso(usuario, dia1, f, hip, 'total', total=D(kg))
    r = cliente.get(reverse('apptrainer:semana_excel'), {'desde': '2026-03-09', 'hasta': '2026-03-15'})
    assert 'registro-2026-03-09_a_2026-03-15.xlsx' in r['Content-Disposition']
    assert [f[11] for f in _filas(r) if f[4] == 'Hip thrust'] == [40.0]
    amplio = cliente.get(reverse('apptrainer:semana_excel'), {'desde': '2026-03-01', 'hasta': '2026-03-31'})
    assert sorted(f[11] for f in _filas(amplio) if f[4] == 'Hip thrust') == [30.0, 40.0, 50.0]


def test_fechas_invertidas_se_ordenan_e_invalidas_usan_la_semana_actual(usuario):
    hoy = dt.date(2026, 10, 14)
    assert services.rango_valido('2026-03-20', '2026-03-10') == (dt.date(2026, 3, 10), dt.date(2026, 3, 20))
    assert services.rango_valido('x', '2026-03-10', hoy) == (dt.date(2026, 10, 12), dt.date(2026, 10, 18))
    assert services.rango_valido(None, None, hoy) == (dt.date(2026, 10, 12), dt.date(2026, 10, 18))


def test_excel_incluye_detalle_e_indicaciones(cliente, dia1, usuario):
    hip = ej(dia1, 'Hip thrust')
    aducc = ej(dia1, 'Máquina de aducción')
    _sesion_con_peso(usuario, dia1, timezone.localdate(), hip, 'barra', barra=D(20), discos=D(30))
    r = cliente.get(reverse('apptrainer:semana_excel'))
    wb = load_workbook(io.BytesIO(r.content))
    encabezados = [c.value for c in wb['Registro'][1]]
    assert encabezados[7:] == ['Tipo de carga', 'Barra (kg)', 'Discos (kg)', 'Por mancuerna (kg)', 'Total (kg)', 'Indicaciones']
    hip_fila = next(f for f in _filas(r) if f[4] == 'Hip thrust')
    assert hip_fila[7:12] == ('Barra + discos', 20.0, 30.0, None, 50.0)
    assert next(f for f in _filas(r) if f[4] == 'Máquina de aducción')[12] == 'Mantengo 3 seg cerrando'
    assert [c.value for c in wb['Resumen'][1]][-1] == 'Carga total (kg)'
    assert _filas(r, 'Resumen')[0][6] == 50.0


def test_semana_tiene_selector_de_fechas(cliente, con_rutina):
    html = cliente.get(reverse('apptrainer:semana')).content.decode()
    assert 'name="desde"' in html and 'name="hasta"' in html
    assert 'Semana pasada' in html and 'Todo el historial' in html


# ---------- indicaciones del entrenador e imágenes ----------

def test_indicaciones_del_entrenador_textuales(con_rutina):
    notas = {e.nombre: e.nota for e in Ejercicio.objects.filter(bloque__dia__rutina__usuario=con_rutina)}
    assert notas['Máquina de aducción'] == 'Mantengo 3 seg cerrando'
    assert notas['Drop serie en máquina de abducción'] == 'Vas bajando el peso'
    assert notas['Bíceps con barra'] == 'Barra normal'
    assert notas['Sentadilla con empuje'] == 'Alternar con máquina sysy (usar mancuernas)'
    assert notas['Remo horizontal en máquina'] == 'Agarre neutro. Superserie'
    assert 'Bíceps martillo iso alternado' in notas and 'Bíceps martillo isométrico alternado' not in notas


def test_las_indicaciones_se_ven_en_la_pantalla_del_dia(cliente, dia1):
    html = cliente.get(reverse('apptrainer:dia', args=[dia1.pk])).content.decode()
    assert 'Mantengo 3 seg cerrando' in html and 'ex-note-label">Entrenador' in html


def test_cada_ejercicio_del_catalogo_tiene_su_imagen_y_descripcion():
    from pathlib import Path
    carpeta = Path(__file__).resolve().parents[1] / 'static' / 'apptrainer' / 'ejercicios'
    for nombre, (imagen, _) in CATALOGO.items():
        assert (carpeta / f'{imagen}.svg').is_file(), f'falta la imagen de {nombre}'
        assert DESCRIPCIONES.get(imagen), f'falta la descripción de {imagen}'


def test_todos_los_ejercicios_sembrados_tienen_imagen(con_rutina):
    assert not Ejercicio.objects.filter(bloque__dia__rutina__usuario=con_rutina, imagen='').exists()


def test_la_pantalla_del_dia_muestra_las_imagenes(cliente, dia1):
    html = cliente.get(reverse('apptrainer:dia', args=[dia1.pk])).content.decode()
    assert html.count('class="ex-img"') == Ejercicio.objects.filter(bloque__dia=dia1).count() == 12
    assert 'ejercicios/hip-thrust.svg' in html and 'alt="Espalda alta en el banco' in html


def test_imagen_inexistente_no_rompe_la_pagina(cliente, dia1):
    Ejercicio.objects.filter(bloque__dia=dia1).update(imagen='no-existe')
    assert cliente.get(reverse('apptrainer:dia', args=[dia1.pk])).status_code == 200


# ---------- seed: no borra el historial por accidente ----------

def test_seed_no_reemplaza_una_rutina_con_historial(cliente, dia1, usuario):
    _sesion_con_peso(usuario, dia1, timezone.localdate(), ej(dia1, 'Hip thrust'), 'total', total=D(40))
    with pytest.raises(CommandError, match='--forzar'):
        call_command('seed_rutina', usuario='grecia')
    assert PesoEjercicio.objects.count() == 1 and Sesion.objects.count() == 1


def test_seed_forzar_reemplaza_y_borra_el_historial(dia1, usuario):
    _sesion_con_peso(usuario, dia1, timezone.localdate(), ej(dia1, 'Hip thrust'), 'total', total=D(40))
    call_command('seed_rutina', usuario='grecia', forzar=True, stdout=open('/dev/null', 'w'))
    assert not Sesion.objects.exists() and Dia.objects.filter(rutina__usuario=usuario).count() == 3
