import io
from decimal import Decimal as D

import pytest
from django.db import connection
from django.db.migrations.executor import MigrationExecutor
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apptrainer import services
from apptrainer.cargas import armar_carga, detalle_texto
from apptrainer.models import Dia, Ejercicio, PesoEjercicio, Sesion

from .test_dia_pesos import dia1, ej  # noqa: F401  (fixtures y helpers)


def guardar(cliente, dia, fecha=None, **campos):
    datos = {'fecha': (fecha or timezone.localdate()).isoformat()}
    datos.update(campos)
    return cliente.post(reverse('apptrainer:dia_guardar', args=[dia.pk]), datos)


def pagina_dia(cliente, dia, **params):
    return cliente.get(reverse('apptrainer:dia', args=[dia.pk]), params).content.decode()


# ---------- hip thrust: barra + cada lado, sin mancuernas ----------

def test_barra_con_peso_por_lado():
    c = armar_carga('barra', barra=D(20), por_lado=D(15), lados=True)
    assert c.kg == 50 and c.barra == 20 and c.por_lado == 15 and c.discos is None
    assert detalle_texto('barra', D(20), None, None, D(50), D(15)) == '20 + 2 × 15 = 50 kg'
    assert armar_carga('barra', por_lado=D(10), lados=True).kg == 20
    assert armar_carga('barra', lados=True) is None


def test_hip_thrust_se_configura_con_barra_y_cada_lado(con_rutina):
    hip = Ejercicio.objects.get(bloque__dia__rutina__usuario=con_rutina, nombre='Hip thrust')
    assert hip.carga_por_lado is True and hip.modos == ['barra', 'total']
    otro = Ejercicio.objects.get(bloque__dia__rutina__usuario=con_rutina, nombre='Peso muerto rumano')
    assert otro.carga_por_lado is False and otro.modos == ['barra', 'mancuernas', 'total']


def test_hip_thrust_no_ofrece_mancuernas_y_pide_cada_lado(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    html = pagina_dia(cliente, dia1)
    assert f'id="modo_{hip.pk}_barra"' in html and f'id="modo_{hip.pk}_total"' in html
    assert f'id="modo_{hip.pk}_mancuernas"' not in html
    assert f'name="lado_{hip.pk}"' in html and f'name="discos_{hip.pk}"' not in html
    # los demás ejercicios siguen igual
    rumano = ej(dia1, 'Peso muerto rumano')
    assert f'id="modo_{rumano.pk}_mancuernas"' in html and f'name="discos_{rumano.pk}"' in html


def test_guardar_hip_thrust_barra_y_cada_lado(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    r = guardar(cliente, dia1, **{f'modo_{hip.pk}': 'barra', f'barra_{hip.pk}': '20', f'lado_{hip.pk}': '15'})
    assert r.status_code == 302
    p = PesoEjercicio.objects.get()
    assert (p.modo, p.kg, p.kg_barra, p.kg_por_lado, p.kg_discos) == ('barra', D(50), D(20), D(15), None)
    assert p.detalle == '20 + 2 × 15 = 50 kg'


def test_hip_thrust_se_vuelve_a_mostrar_con_lo_guardado(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    guardar(cliente, dia1, **{f'modo_{hip.pk}': 'barra', f'barra_{hip.pk}': '20', f'lado_{hip.pk}': '12,5'})
    html = pagina_dia(cliente, dia1)
    assert 'value="12,5"' in html and 'Total: <b data-total>45</b>' in html


def test_hip_thrust_valida_cada_lado(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    r = guardar(cliente, dia1, **{f'modo_{hip.pk}': 'barra', f'barra_{hip.pk}': '20', f'lado_{hip.pk}': 'abc'})
    assert r.status_code == 422 and not PesoEjercicio.objects.exists()
    assert 'value="20"' in r.content.decode() and 'value="abc"' in r.content.decode()


def test_hip_thrust_ignora_mancuernas_enviadas(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    guardar(cliente, dia1, **{f'modo_{hip.pk}': 'mancuernas', f'mancuerna_{hip.pk}': '12', f'kg_{hip.pk}': '40'})
    p = PesoEjercicio.objects.get()
    assert p.modo == 'total' and p.kg == 40 and p.kg_mancuerna is None


def test_hip_thrust_total_maximo(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    r = guardar(cliente, dia1, **{f'modo_{hip.pk}': 'barra', f'barra_{hip.pk}': '20', f'lado_{hip.pk}': '300'})
    assert r.status_code == 422 and 'no puede superar 500' in r.content.decode()


def test_registro_viejo_de_hip_thrust_sigue_viéndose(cliente, dia1, usuario):
    """Lo guardado antes del cambio (barra + discos, o mancuernas) no se pierde ni rompe la pantalla."""
    hip = ej(dia1, 'Hip thrust')
    ayer = timezone.localdate() - __import__('datetime').timedelta(days=3)
    services.guardar_pesos(usuario, dia1, ayer, {hip.pk: armar_carga('barra', barra=D(20), discos=D(30))})
    html = pagina_dia(cliente, dia1)
    assert 'Anterior: 20 + 30 = 50 kg' in html
    services.guardar_pesos(usuario, dia1, ayer, {hip.pk: armar_carga('mancuernas', mancuerna=D(10))})
    html = pagina_dia(cliente, dia1, fecha=ayer.isoformat())
    assert f'id="modo_{hip.pk}_barra" name="modo_{hip.pk}" value="barra" form="form-guardar" checked' in html
    assert PesoEjercicio.objects.get().detalle == '2 × 10 = 20 kg'


def test_excel_incluye_cada_lado(cliente, dia1, usuario):
    hip = ej(dia1, 'Hip thrust')
    guardar(cliente, dia1, **{f'modo_{hip.pk}': 'barra', f'barra_{hip.pk}': '20', f'lado_{hip.pk}': '15'})
    filas = list(load_workbook(io.BytesIO(cliente.get(reverse('apptrainer:semana_excel')).content))['Registro'].iter_rows(min_row=2, values_only=True))
    fila = next(f for f in filas if f[4] == 'Hip thrust')
    assert fila[7:13] == ('Barra + discos', 20.0, None, 15.0, None, 50.0)


# ---------- comentarios de la rutina ----------

def test_la_pantalla_del_dia_tiene_la_caja_de_comentarios(cliente, dia1):
    html = pagina_dia(cliente, dia1)
    assert '<textarea id="comentario" name="comentario" form="form-guardar"' in html
    assert 'Comentarios de la rutina' in html and 'maxlength="1000"' in html


def test_guardar_comentario_junto_con_los_pesos(cliente, dia1, usuario):
    hip = ej(dia1, 'Hip thrust')
    guardar(cliente, dia1, comentario='  Me sentí fuerte hoy.  ', **{f'kg_{hip.pk}': '40'})
    s = Sesion.objects.get(usuario=usuario)
    assert s.comentario == 'Me sentí fuerte hoy.'
    assert 'Me sentí fuerte hoy.' in pagina_dia(cliente, dia1)


def test_se_puede_guardar_solo_un_comentario(cliente, dia1, usuario):
    r = guardar(cliente, dia1, comentario='Hoy no pude hacer todo')
    assert r.status_code == 302
    assert Sesion.objects.get(usuario=usuario).comentario == 'Hoy no pude hacer todo'
    assert not PesoEjercicio.objects.exists()


def test_comentario_vacio_sin_pesos_no_crea_sesion(cliente, dia1):
    guardar(cliente, dia1, comentario='   ')
    assert not Sesion.objects.exists()


def test_el_comentario_se_edita_y_se_borra(cliente, dia1, usuario):
    guardar(cliente, dia1, comentario='uno')
    guardar(cliente, dia1, comentario='dos')
    assert Sesion.objects.get().comentario == 'dos'
    guardar(cliente, dia1, comentario='')
    assert Sesion.objects.get().comentario == ''


def test_guardar_pesos_sin_comentario_no_lo_borra(dia1, usuario):
    hip = ej(dia1, 'Hip thrust')
    hoy = timezone.localdate()
    services.guardar_pesos(usuario, dia1, hoy, {}, comentario='importante')
    services.guardar_pesos(usuario, dia1, hoy, {hip.pk: armar_carga('total', total=D(40))})  # sin comentario=
    assert Sesion.objects.get().comentario == 'importante'


def test_comentario_demasiado_largo_no_guarda_nada_y_conserva_el_texto(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    largo = 'a' * 1001
    r = guardar(cliente, dia1, comentario=largo, **{f'kg_{hip.pk}': '40'})
    assert r.status_code == 422 and 'muy largo' in r.content.decode()
    assert not Sesion.objects.exists() and not PesoEjercicio.objects.exists()
    assert largo in r.content.decode()


def test_error_de_peso_no_pierde_el_comentario(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    r = guardar(cliente, dia1, comentario='mi nota', **{f'kg_{hip.pk}': 'abc'})
    assert r.status_code == 422 and 'mi nota' in r.content.decode() and not Sesion.objects.exists()


def test_comentarios_en_resumen_cierre_y_excel(cliente, dia1, usuario):
    import datetime as dt
    hoy = timezone.localdate()
    hip = ej(dia1, 'Hip thrust')
    pasada = services.lunes_de(hoy) - dt.timedelta(days=4)
    services.guardar_pesos(usuario, dia1, pasada, {hip.pk: armar_carga('total', total=D(40))}, comentario='Semana pasada: dolor de rodilla')
    guardar(cliente, dia1, comentario='Esta semana mejor\ncon calentamiento', **{f'kg_{hip.pk}': '45'})

    resumen = cliente.get(reverse('apptrainer:dia_resumen', args=[dia1.pk])).content.decode()
    assert 'Tus comentarios' in resumen and 'Esta semana mejor<br>con calentamiento' in resumen

    home = cliente.get(reverse('apptrainer:hoy')).content.decode()
    assert 'Semana pasada: dolor de rodilla' in home and 'Esta semana mejor' in home

    wb = load_workbook(io.BytesIO(cliente.get(reverse('apptrainer:semana_excel'), {'todo': '1'}).content))
    comentarios = [f[7] for f in wb['Resumen'].iter_rows(min_row=2, values_only=True)]
    assert 'Semana pasada: dolor de rodilla' in comentarios and 'Esta semana mejor\ncon calentamiento' in comentarios


def test_comentarios_son_privados(client, dia1, usuario, otro_usuario):
    services.guardar_pesos(usuario, dia1, timezone.localdate(), {}, comentario='secreto')
    client.force_login(otro_usuario)
    assert client.post(reverse('apptrainer:dia_guardar', args=[dia1.pk]), {'comentario': 'hack'}).status_code == 404
    assert client.get(reverse('apptrainer:dia_resumen', args=[dia1.pk])).status_code == 404
    assert 'secreto' not in client.get(reverse('apptrainer:hoy')).content.decode()
    assert Sesion.objects.get().comentario == 'secreto'


# ---------- las migraciones no pierden historial ----------

@pytest.mark.django_db(transaction=True)
def test_aplicar_migraciones_nuevas_conserva_todo_el_historial():
    """Simula la base de producción antes de este cambio (en 0005), con historial real, aplica
    las migraciones nuevas y comprueba que nada se perdió ni cambió."""
    antes = [('apptrainer', '0005_indicaciones_textuales_del_entrenador')]
    ejecutor = MigrationExecutor(connection)
    ejecutor.migrate(antes)
    try:
        apps = MigrationExecutor(connection).loader.project_state(antes).apps
        User, Rutina, Dia, Bloque = (apps.get_model('auth', 'User'), apps.get_model('apptrainer', 'Rutina'),
                                     apps.get_model('apptrainer', 'Dia'), apps.get_model('apptrainer', 'Bloque'))
        Ejercicio, Sesion_, Hecho, Peso, Corporal = (apps.get_model('apptrainer', n) for n in
                                                      ('Ejercicio', 'Sesion', 'EjercicioHecho', 'PesoEjercicio', 'RegistroPeso'))
        u = User.objects.create(username='grecia')
        rutina = Rutina.objects.create(usuario=u, nombre='Rutina semanal', activa=True)
        dia = Dia.objects.create(rutina=rutina, orden=1, titulo='Día 1')
        bloque = Bloque.objects.create(dia=dia, tipo='principal', orden=1)
        hip = Ejercicio.objects.create(bloque=bloque, orden=1, nombre='Hip thrust', nota='Barra normal')
        rumano = Ejercicio.objects.create(bloque=bloque, orden=2, nombre='Peso muerto rumano')
        sesiones = [Sesion_.objects.create(usuario=u, dia=dia, fecha=f'2026-10-{d:02d}', completada=bool(d % 2)) for d in (5, 12, 19)]
        for s in sesiones:
            Hecho.objects.create(sesion=s, ejercicio=hip)
            Peso.objects.create(sesion=s, ejercicio=hip, modo='barra', kg_barra=20, kg_discos=30, kg=50)
            Peso.objects.create(sesion=s, ejercicio=rumano, modo='mancuernas', kg_mancuerna=12, kg=24)
        Corporal.objects.create(usuario=u, fecha='2026-10-01', kg=63, nota='primer registro')
        foto_antes = {
            'sesiones': sorted(Sesion_.objects.values_list('fecha', 'completada')),
            'hechos': Hecho.objects.count(),
            'pesos': sorted(Peso.objects.values_list('sesion__fecha', 'ejercicio__nombre', 'modo', 'kg_barra', 'kg_discos', 'kg_mancuerna', 'kg')),
            'corporal': list(Corporal.objects.values_list('fecha', 'kg', 'nota')),
        }

        ultimo = MigrationExecutor(connection)
        ultimo.migrate(ultimo.loader.graph.leaf_nodes())

        from apptrainer import models as m
        assert sorted(m.Sesion.objects.values_list('fecha', 'completada')) == foto_antes['sesiones']
        assert m.EjercicioHecho.objects.count() == foto_antes['hechos']
        assert sorted(m.PesoEjercicio.objects.values_list('sesion__fecha', 'ejercicio__nombre', 'modo', 'kg_barra', 'kg_discos', 'kg_mancuerna', 'kg')) == foto_antes['pesos']
        assert list(m.RegistroPeso.objects.values_list('fecha', 'kg', 'nota')) == foto_antes['corporal']
        # lo nuevo arranca vacío y el hip thrust se reconfigura sin tocar sus registros
        assert not m.Sesion.objects.exclude(comentario='').exists()
        assert not m.PesoEjercicio.objects.exclude(kg_por_lado=None).exists()
        hip = m.Ejercicio.objects.get(nombre='Hip thrust')
        assert hip.modos == ['barra', 'total'] and hip.carga_por_lado is True and hip.nota == 'Barra normal'
        assert m.PesoEjercicio.objects.filter(ejercicio=hip).first().detalle == '20 + 30 = 50 kg'
    finally:
        final = MigrationExecutor(connection)
        final.migrate(final.loader.graph.leaf_nodes())
