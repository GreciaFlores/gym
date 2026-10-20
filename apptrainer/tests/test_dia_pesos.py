import datetime as dt
import io
from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apptrainer import services
from apptrainer.models import Dia, Ejercicio, EjercicioHecho, PesoEjercicio, Sesion


@pytest.fixture
def dia1(con_rutina):
    return Dia.objects.get(rutina__usuario=con_rutina, orden=1)


def ej(dia, nombre):
    return Ejercicio.objects.get(bloque__dia=dia, nombre=nombre)


def manana():
    return timezone.localdate() + dt.timedelta(days=1)


def guardar(cliente, dia, fecha=None, **kg):
    datos = {'fecha': (fecha or timezone.localdate()).isoformat()}
    datos.update({f'kg_{k}': v for k, v in kg.items()})
    return cliente.post(reverse('apptrainer:dia_guardar', args=[dia.pk]), datos)


# ---------- fecha de la sesión ----------

def test_dia_usa_la_fecha_elegida(cliente, dia1):
    html = cliente.get(reverse('apptrainer:dia', args=[dia1.pk]), {'fecha': manana().isoformat()}).content.decode()
    assert f'value="{manana().isoformat()}"' in html
    assert 'type="date"' in html and 'Mañana' in html


@pytest.mark.parametrize('basura', ['', 'hola', '2026-13-45', '1900-01-01', '2999-01-01'])
def test_fecha_invalida_o_absurda_usa_hoy(cliente, dia1, basura):
    html = cliente.get(reverse('apptrainer:dia', args=[dia1.pk]), {'fecha': basura}).content.decode()
    assert f'value="{timezone.localdate().isoformat()}"' in html


def test_marcar_ejercicio_en_otra_fecha_crea_la_sesion_de_esa_fecha(cliente, dia1, usuario):
    e = ej(dia1, 'Hip thrust')
    cliente.post(reverse('apptrainer:alternar', args=[e.pk]), {'fecha': manana().isoformat()},
                 HTTP_X_REQUESTED_WITH='fetch')
    assert Sesion.objects.get(usuario=usuario, dia=dia1).fecha == manana()
    assert EjercicioHecho.objects.filter(sesion__fecha=manana(), ejercicio=e).exists()
    # y no aparece en hoy
    assert not EjercicioHecho.objects.filter(sesion__fecha=timezone.localdate()).exists()


# ---------- pesos por ejercicio ----------

def test_pesos_solo_en_zona_media_y_principal(cliente, dia1):
    html = cliente.get(reverse('apptrainer:dia', args=[dia1.pk])).content.decode()
    ids_con_peso = Ejercicio.objects.filter(bloque__dia=dia1).exclude(bloque__tipo='movilidad')
    ids_movilidad = Ejercicio.objects.filter(bloque__dia=dia1, bloque__tipo='movilidad')
    assert all(f'name="kg_{e.pk}"' in html for e in ids_con_peso)
    assert not any(f'name="kg_{e.pk}"' in html for e in ids_movilidad)
    assert html.count('data-carga') == ids_con_peso.count() == 8


def test_guardar_pesos_con_coma_y_punto(cliente, dia1, usuario):
    r = guardar(cliente, dia1, **{str(ej(dia1, 'Hip thrust').pk): '40,5',
                                  str(ej(dia1, 'Peso muerto rumano').pk): '25.25'})
    assert r.status_code == 302
    assert r['Location'] == f"{reverse('apptrainer:dia_resumen', args=[dia1.pk])}?fecha={timezone.localdate().isoformat()}"
    pesos = {p.ejercicio.nombre: p.kg for p in PesoEjercicio.objects.all()}
    assert pesos == {'Hip thrust': Decimal('40.50'), 'Peso muerto rumano': Decimal('25.25')}


def test_guardar_actualiza_y_borra_al_vaciar(cliente, dia1):
    hip, rumano = ej(dia1, 'Hip thrust'), ej(dia1, 'Peso muerto rumano')
    guardar(cliente, dia1, **{str(hip.pk): '40', str(rumano.pk): '20'})
    guardar(cliente, dia1, **{str(hip.pk): '42,5', str(rumano.pk): ''})
    assert PesoEjercicio.objects.get().kg == Decimal('42.50')
    assert PesoEjercicio.objects.count() == 1  # el vaciado borró el de rumano


def test_guardar_sin_pesos_no_crea_sesion_vacia(cliente, dia1):
    guardar(cliente, dia1, **{str(ej(dia1, 'Hip thrust').pk): ''})
    assert not Sesion.objects.exists()


@pytest.mark.parametrize('malo', ['abc', '0', '-5', '501', '1000'])
def test_peso_invalido_no_guarda_nada_y_muestra_lo_escrito(cliente, dia1, malo):
    hip, rumano = ej(dia1, 'Hip thrust'), ej(dia1, 'Peso muerto rumano')
    r = guardar(cliente, dia1, **{str(hip.pk): '40', str(rumano.pk): malo})
    html = r.content.decode()
    assert r.status_code == 422
    assert not PesoEjercicio.objects.exists()
    assert 'field-error' in html and 'no se guardó nada' in html
    assert f'value="{malo}"' in html and 'value="40"' in html  # conserva lo tipeado


def test_kg_de_movilidad_se_ignora(cliente, dia1):
    movilidad = Ejercicio.objects.filter(bloque__dia=dia1, bloque__tipo='movilidad').first()
    guardar(cliente, dia1, **{str(movilidad.pk): '30', str(ej(dia1, 'Hip thrust').pk): '40'})
    assert list(PesoEjercicio.objects.values_list('ejercicio__nombre', flat=True)) == ['Hip thrust']


def test_guardar_en_otra_fecha(cliente, dia1, usuario):
    guardar(cliente, dia1, fecha=manana(), **{str(ej(dia1, 'Hip thrust').pk): '40'})
    assert Sesion.objects.get(usuario=usuario).fecha == manana()


def test_peso_anterior_aparece_como_referencia(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    ayer = timezone.localdate() - dt.timedelta(days=3)
    guardar(cliente, dia1, fecha=ayer, **{str(hip.pk): '37,5'})
    html = cliente.get(reverse('apptrainer:dia', args=[dia1.pk])).content.decode()
    assert 'Anterior: 37,5 kg' in html
    # el valor de hoy no se toma como "anterior" de sí mismo
    guardar(cliente, dia1, **{str(hip.pk): '40'})
    html = cliente.get(reverse('apptrainer:dia', args=[dia1.pk])).content.decode()
    assert 'Anterior: 37,5 kg' in html and 'value="40"' in html


def test_guardar_exige_post_y_login(client, cliente, dia1):
    assert cliente.get(reverse('apptrainer:dia_guardar', args=[dia1.pk])).status_code == 405
    client.logout()
    r = client.post(reverse('apptrainer:dia_guardar', args=[dia1.pk]))
    assert r.status_code == 302 and '/entrar/' in r['Location']


# ---------- resumen ----------

def test_resumen_muestra_hechos_y_pesos(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    cliente.post(reverse('apptrainer:alternar', args=[hip.pk]), HTTP_X_REQUESTED_WITH='fetch')
    guardar(cliente, dia1, **{str(hip.pk): '42,5'})
    r = cliente.get(reverse('apptrainer:dia_resumen', args=[dia1.pk]))
    html = r.content.decode()
    assert 'Resumen · Día 1' in html and '42,5 kg' in html
    assert '1/12' in html and '1/8' in html  # hechos y pesos anotados


def test_resumen_vacio(cliente, dia1):
    html = cliente.get(reverse('apptrainer:dia_resumen', args=[dia1.pk])).content.decode()
    assert 'Todavía no hay registro' in html


# ---------- aislamiento ----------

def test_otro_usuario_no_puede_guardar_ni_ver_resumen(client, dia1, otro_usuario):
    client.force_login(otro_usuario)
    assert client.post(reverse('apptrainer:dia_guardar', args=[dia1.pk]), {'fecha': '2026-01-01'}).status_code == 404
    assert client.get(reverse('apptrainer:dia_resumen', args=[dia1.pk])).status_code == 404
    assert not PesoEjercicio.objects.exists()


# ---------- Excel semanal ----------

def _abrir(respuesta):
    return load_workbook(io.BytesIO(respuesta.content))


def test_semana_tiene_boton_de_excel(cliente, con_rutina):
    html = cliente.get(reverse('apptrainer:semana')).content.decode()
    assert reverse('apptrainer:semana_excel') in html and 'Exportar Excel' in html


def test_excel_semanal_incluye_hechos_y_pesos(cliente, dia1):
    hip = ej(dia1, 'Hip thrust')
    cliente.post(reverse('apptrainer:alternar', args=[hip.pk]), HTTP_X_REQUESTED_WITH='fetch')
    guardar(cliente, dia1, **{str(hip.pk): '42,5'})
    r = cliente.get(reverse('apptrainer:semana_excel'))
    assert r.status_code == 200 and r['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'attachment; filename="apptrainer-registro-' in r['Content-Disposition']
    wb = _abrir(r)
    filas = list(wb['Registro'].iter_rows(min_row=2, values_only=True))
    assert len(filas) == 12  # todos los ejercicios del día
    hip_fila = next(f for f in filas if f[4] == 'Hip thrust')
    assert hip_fila[1] == 'Día 1' and hip_fila[3] == 'Trabajo principal' and hip_fila[6] == 'Sí'
    assert hip_fila[7] == 'Total' and hip_fila[11] == 42.5
    rumano = next(f for f in filas if f[4] == 'Peso muerto rumano')
    assert rumano[6:12] == ('No', None, None, None, None, None)
    resumen = list(wb['Resumen'].iter_rows(min_row=2, values_only=True))
    assert resumen[0][3:] == ('En curso', 1, 12, 42.5)


def test_excel_semanal_solo_la_semana_actual_y_todo_el_historial(cliente, dia1, usuario):
    hip = ej(dia1, 'Hip thrust')
    vieja = timezone.localdate() - dt.timedelta(days=21)
    guardar(cliente, dia1, fecha=vieja, **{str(hip.pk): '30'})
    guardar(cliente, dia1, **{str(hip.pk): '40'})
    semana = _abrir(cliente.get(reverse('apptrainer:semana_excel')))
    assert [f[11] for f in semana['Registro'].iter_rows(min_row=2, values_only=True) if f[4] == 'Hip thrust'] == [40.0]
    todo = _abrir(cliente.get(reverse('apptrainer:semana_excel'), {'todo': '1'}))
    assert sorted(f[11] for f in todo['Registro'].iter_rows(min_row=2, values_only=True) if f[4] == 'Hip thrust') == [30.0, 40.0]
    assert 'historial' in cliente.get(reverse('apptrainer:semana_excel'), {'todo': '1'})['Content-Disposition']


def test_excel_semanal_vacio_no_falla(cliente, con_rutina):
    wb = _abrir(cliente.get(reverse('apptrainer:semana_excel')))
    assert wb['Registro'].max_row == 1 and wb['Resumen'].max_row == 1


def test_excel_semanal_solo_datos_propios(client, dia1, usuario, otro_usuario):
    from apptrainer.cargas import armar_carga
    services.guardar_pesos(usuario, dia1, timezone.localdate(),
                           {ej(dia1, 'Hip thrust').pk: armar_carga('total', total=Decimal('40'))})
    client.force_login(otro_usuario)
    wb = _abrir(client.get(reverse('apptrainer:semana_excel')))
    assert wb['Registro'].max_row == 1


def test_excel_exige_login(client, db):
    r = client.get(reverse('apptrainer:semana_excel'))
    assert r.status_code == 302 and '/entrar/' in r['Location']
